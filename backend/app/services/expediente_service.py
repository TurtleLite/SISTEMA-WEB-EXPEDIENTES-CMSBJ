from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Color
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session
import datetime
import math
from app.models.list_definition import ListDefinition, ListRecord
from app.schemas.list_definition import ListDefinitionCreate



EXPEDIENTE_COLUMNS = [
    {"key": "especialidad", "label": "Especialidad", "type": "text"},
    {"key": "criticidad", "label": "Criticidad Clínica", "type": "text"},
    {"key": "estatus", "label": "Estatus del Paciente", "type": "text"},
    {"key": "nombre", "label": "Nombre/First Name", "type": "text"},
    {"key": "apellido", "label": "Apellido/Last Name", "type": "text"},
    {"key": "sexo", "label": "Sexo/Sex", "type": "text"},
    {"key": "edad", "label": "Age/Edad", "type": "number"},
    {"key": "fecha_elaboracion", "label": "Fecha de Elaboración", "type": "date"},
    {"key": "identidad", "label": "Nº Identidad", "type": "text"},
    {"key": "persona_responsable", "label": "Persona Responsable", "type": "text"},
    {"key": "albergue", "label": "Albergue", "type": "text"},
    {"key": "perfil", "label": "Perfil", "type": "text"},
    {"key": "telefono", "label": "Teléfono", "type": "text"},
    {"key": "telefono2", "label": "Teléfono 2", "type": "text"},
    {"key": "telefono3", "label": "Teléfono 3", "type": "text"},
    {"key": "expediente", "label": "Expediente", "type": "text"},
    {"key": "domicilio", "label": "Domicilio del Paciente", "type": "text"},
    {"key": "historia_enfermedad", "label": "Historia de Enfermedad Actual", "type": "text"},
    {"key": "enfermedades_previas", "label": "Enfermedades Anteriores", "type": "text"},
    {"key": "cirugias_previas", "label": "Cirugías Anteriores", "type": "text"},
    {"key": "alergias", "label": "Alergias", "type": "text"},
    {"key": "otros_antecedentes", "label": "Otros Antecedentes", "type": "text"},
    {"key": "presion_arterial", "label": "P.A./B.P", "type": "text"},
    {"key": "fc", "label": "F.C.", "type": "text"},
    {"key": "pulso", "label": "Pulso", "type": "text"},
    {"key": "temperatura", "label": "T°", "type": "text"},
    {"key": "fr", "label": "F.R.", "type": "text"},
    {"key": "peso", "label": "Peso/Weight", "type": "text"},
    {"key": "talla", "label": "Talla", "type": "text"},
    {"key": "bmi", "label": "B.M.I.", "type": "text"},
    {"key": "examen_fisico", "label": "Examen Físico", "type": "text"},
    {"key": "diagnostico", "label": "Diagnóstico", "type": "text"},
    {"key": "nombre_medico", "label": "Nombre del Médico", "type": "text"},
    {"key": "estatus_cirugia", "label": "Estatus de Cirugía", "type": "text"},
]


def create_expediente_template(db: Session, user_id: int) -> ListDefinition:
    from app.services.list_service import create_list_definition
    schema = ListDefinitionCreate(
        name="Expediente Médico",
        description="Historial clínico de pacientes con datos personales, antecedentes, signos vitales y diagnóstico",
        columns_config=EXPEDIENTE_COLUMNS,
    )
    return create_list_definition(db, schema, user_id)


def _thin_border():
    return Border(
        left=Side(style='thin', color=Color(auto=True)),
        right=Side(style='thin', color=Color(auto=True)),
        top=Side(style='thin', color=Color(auto=True)),
        bottom=Side(style='thin', color=Color(auto=True)),
    )

def _medium_bottom():
    return Border(bottom=Side(style='medium', color=Color(auto=True)))

def _medium_top():
    return Border(top=Side(style='medium', color=Color(auto=True)))

def _thin_border_right_medium():
    return Border(left=Side(style='thin', color=Color(auto=True)), right=Side(style='medium', color=Color(auto=True)), top=Side(style='thin', color=Color(auto=True)), bottom=Side(style='thin', color=Color(auto=True)))

def _thin_border_left_medium():
    return Border(left=Side(style='medium', color=Color(auto=True)), top=Side(style='thin', color=Color(auto=True)), bottom=Side(style='thin', color=Color(auto=True)))

def _thin_border_bottom_only():
    return Border(bottom=Side(style='thin', color=Color(auto=True)))

def _thin_border_top_bottom():
    return Border(top=Side(style='thin', color=Color(auto=True)), bottom=Side(style='thin', color=Color(auto=True)))

def _thin_border_top():
    return Border(top=Side(style='thin', color=Color(auto=True)))

def _thin_border_bottom():
    return Border(bottom=Side(style='thin', color=Color(auto=True)))

def _thin_border_sides():
    return Border(left=Side(style='thin', color=Color(auto=True)), right=Side(style='thin', color=Color(auto=True)))

def _thin_border_left():
    return Border(left=Side(style='thin', color=Color(auto=True)))

def _thin_border_right():
    return Border(right=Side(style='thin', color=Color(auto=True)))


def export_expediente_excel(records: list[ListRecord], filepath: str, logo_path: str = None):
    import os
    from openpyxl.drawing.image import Image

    wb = Workbook()
    wb._fonts[0] = Font(name='Arial', size=10)
    wb._named_styles['Normal'].font = Font(name='Arial', size=10)

    arial = 'Arial'
    center_wrap = Alignment(horizontal="center", vertical="center", wrap_text=True)
    center_nowrap = Alignment(horizontal="center", vertical="center")
    center_vwrap = Alignment(vertical="center", wrap_text=True)
    center_hnone_wrap = Alignment(horizontal="center", wrap_text=True)
    center_hnone = Alignment(horizontal="center")
    left_center_wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)
    left_center_nowrap = Alignment(horizontal="left", vertical="center")
    right_center_wrap = Alignment(horizontal="right", vertical="center", wrap_text=True)
    center_h_wrap = Alignment(horizontal="center", wrap_text=True)
    center_h = Alignment(horizontal="center")
    vtop = Alignment(vertical="top")

    thin = _thin_border()
    thin_top_bottom = _thin_border_top_bottom()
    thin_top = _thin_border_top()
    thin_bottom = _thin_border_bottom()
    thin_sides = _thin_border_sides()
    thin_left = _thin_border_left()
    thin_right = _thin_border_right()
    thin_right_medium = _thin_border_right_medium()
    thin_left_medium = _thin_border_left_medium()
    medium_bottom = _medium_bottom()
    medium_top = _medium_top()
    styles = (thin, thin_top_bottom, thin_top, thin_bottom, thin_sides, thin_left, thin_right, thin_right_medium, thin_left_medium, medium_bottom, medium_top)

    col_widths = {
        'A': 11.42578125, 'B': 16.0, 'C': 11.42578125, 'D': 16.0,
        'E': 9.42578125, 'F': 10.140625,         'G': 13.1, 'H': 14.85546875,
        'I': 15.28515625, 'J': 13.7109375, 'K': 11.42578125,
    }

    row_heights = {
        1: 68.0, 2: 24.0, 3: 26.25, 5: 12.75, 6: 12.75,
        8: 12.75, 9: 12.75, 10: 12.75, 14: 13.5, 15: 12.75,
        16: 12.75, 21: 12.75, 25: 12.75, 26: 12.75, 27: 12.75,
        29: 12.75, 30: 12.75, 38: 12.75, 44: 16.5, 48: 13.5, 50: 13.5,
    }

    if logo_path is None:
        _try_paths = [
            os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'assets', 'logo_sbj.png'),
            os.path.join(os.getcwd(), 'backend', 'app', 'assets', 'logo_sbj.png'),
            os.path.join(os.getcwd(), 'app', 'assets', 'logo_sbj.png'),
        ]
        logo_path = None
        for p in _try_paths:
            if os.path.exists(p):
                logo_path = p
                break

    # Remove default empty sheet
    default_sheet = wb.active

    def _write_record_sheet(ws, d, styles):
        ws.page_setup.orientation = 'portrait'
        ws.page_setup.paperSize = 1
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_margins.left = 0.7874015748031495
        ws.page_margins.right = 0.39370078740157477
        ws.page_margins.top = 0.39370078740157477
        ws.page_margins.bottom = 0.19685039370078738
        ws.page_margins.header = 0
        ws.page_margins.footer = 0
        thin, thin_top_bottom, thin_top, thin_bottom, thin_sides, thin_left, thin_right, thin_right_medium, thin_left_medium, medium_bottom, medium_top = styles

        def cellb(r, c, border):
            _apply_border(ws, r, c, border)

        def thin_edges(r1, r2):
            for rr in range(r1 + 1, r2):
                cellb(rr, 1, thin_left)
                cellb(rr, 8, thin_right)

        r = 1

        # === ROW 1: Title + Logo ===
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        c = ws.cell(r, 1, "Centro Médico San Benito José")
        c.font = Font(name=arial, bold=True, size=16, u='single')
        c.alignment = center_wrap
        if logo_path and os.path.exists(logo_path):
            from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
            from openpyxl.drawing.xdr import XDRPositiveSize2D
            from openpyxl.utils.units import pixels_to_EMU
            img = Image(logo_path)
            ratio = img.height / img.width
            img.width = 150
            img.height = round(img.width * ratio)
            img.anchor = OneCellAnchor(
                _from=AnchorMarker(col=0, colOff=0, row=0, rowOff=0),
                ext=XDRPositiveSize2D(cx=pixels_to_EMU(img.width), cy=pixels_to_EMU(img.height)),
            )
            ws.add_image(img)
        r += 1

        # === ROW 2: Especialidad (G-H, size 16) ===
        ws.cell(r, 1).font = Font(name=arial, bold=True, size=11)
        ws.merge_cells(start_row=r, start_column=7, end_row=r, end_column=8)
        c = ws.cell(r, 7, d.get("especialidad", ""))
        c.font = Font(name=arial, bold=True, size=16, color="FF0000")
        c.alignment = Alignment(horizontal="center", vertical="center", shrink_to_fit=True)
        cellb(r, 7, thin)
        cellb(r, 8, Border(right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin')))
        r += 1

        # === ROW 3: Labels ===
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        c = ws.cell(r, 1, "Nombre/First Name")
        c.font = Font(name=arial, bold=True, size=10)
        c.alignment = center_nowrap
        c.border = Border(left=Side(style='thin'), right=Side(style='medium'), top=Side(style='thin'), bottom=Side(style='thin'))
        ws.cell(r, 2).border = Border(right=Side(style='medium'), top=Side(style='thin'), bottom=Side(style='thin'))

        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
        c = ws.cell(r, 3, "Apellido/ Last Name")
        c.font = Font(name=arial, bold=True, size=10)
        c.alignment = center_nowrap
        ws.cell(r, 3).border = thin_left_medium
        ws.cell(r, 4).border = Border(top=Side(style='thin'), bottom=Side(style='thin'))

        c = ws.cell(r, 5, "Sexo/Sex")
        c.font = Font(name=arial, bold=True, size=10)
        c.alignment = center_nowrap
        ws.cell(r, 5).border = thin

        c = ws.cell(r, 6, "Age/Edad")
        c.font = Font(name=arial, bold=True, size=10)
        c.alignment = center_nowrap
        ws.cell(r, 6).border = Border(right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

        ws.merge_cells(start_row=r, start_column=7, end_row=r, end_column=8)
        c = ws.cell(r, 7, "Fecha de Elaboración (d/m/a)")
        c.font = Font(name=arial, bold=True, size=10)
        c.alignment = center_wrap
        cellb(r, 7, thin)
        cellb(r, 8, Border(right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin')))
        r += 1

        # === ROW 4-5: Values ===
        ws.merge_cells(start_row=r, start_column=1, end_row=r+1, end_column=2)
        c = ws.cell(r, 1, d.get("nombre", ""))
        c.font = Font(name=arial, size=14)
        c.alignment = center_wrap
        cellb(r, 1, Border(left=Side(style='thin'), right=Side(style='thin'), bottom=Side(style='thin')))
        cellb(r, 2, Border(right=Side(style='thin')))
        cellb(r+1, 1, Border(left=Side(style='thin'), bottom=Side(style='thin')))
        cellb(r+1, 2, Border(right=Side(style='thin'), bottom=Side(style='thin')))

        ws.merge_cells(start_row=r, start_column=3, end_row=r+1, end_column=4)
        c = ws.cell(r, 3, d.get("apellido", ""))
        c.font = Font(name=arial, size=15)
        c.alignment = center_wrap
        cellb(r, 3, Border(left=Side(style='thin'), right=Side(style='thin'), bottom=Side(style='thin')))
        cellb(r, 4, Border(right=Side(style='thin')))
        cellb(r+1, 3, Border(left=Side(style='thin'), bottom=Side(style='thin')))
        cellb(r+1, 4, Border(right=Side(style='thin'), bottom=Side(style='thin')))

        ws.merge_cells(start_row=r, start_column=5, end_row=r+1, end_column=5)
        c = ws.cell(r, 5, d.get("sexo", ""))
        c.font = Font(name=arial, size=16)
        c.alignment = center_wrap
        cellb(r, 5, Border(left=Side(style='thin'), bottom=Side(style='thin')))
        cellb(r+1, 5, Border(left=Side(style='thin'), bottom=Side(style='thin')))

        ws.merge_cells(start_row=r, start_column=6, end_row=r+1, end_column=6)
        edad_val = d.get("edad", "")
        edad_cell_val = int(edad_val) if isinstance(edad_val, str) and str(edad_val).isdigit() else edad_val
        c = ws.cell(r, 6, edad_cell_val)
        c.font = Font(name=arial, bold=True, size=16)
        c.alignment = center_wrap
        cellb(r, 6, Border(left=Side(style='thin'), bottom=Side(style='thin')))
        cellb(r+1, 6, Border(left=Side(style='thin'), bottom=Side(style='thin')))

        ws.merge_cells(start_row=r, start_column=7, end_row=r+1, end_column=8)
        fecha_raw = d.get("fecha_elaboracion", "")
        fecha_val = fecha_raw
        if isinstance(fecha_raw, str) and fecha_raw.strip():
            try:
                fecha_val = datetime.date.fromisoformat(fecha_raw.strip()[:10])
            except (ValueError, TypeError):
                fecha_val = fecha_raw
        c = ws.cell(r, 7, fecha_val)
        c.font = Font(name=arial, size=16)
        c.alignment = center_wrap
        if isinstance(fecha_val, datetime.date):
            c.number_format = 'mm-dd-yy'
        cellb(r, 7, Border(left=Side(style='thin'), right=Side(style='thin'), bottom=Side(style='thin')))
        cellb(r, 8, Border(right=Side(style='thin')))
        cellb(r+1, 7, Border(left=Side(style='thin'), bottom=Side(style='thin')))
        cellb(r+1, 8, Border(right=Side(style='thin'), bottom=Side(style='thin')))
        r += 2

        # === ROW 6: Labels ===
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        c = ws.cell(r, 1, "Nº Identidad")
        c.font = Font(name=arial, bold=True, size=10)
        c.alignment = center_wrap
        cellb(r, 1, Border(left=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin')))
        cellb(r, 2, Border(top=Side(style='thin'), bottom=Side(style='thin')))

        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4)
        c = ws.cell(r, 3, "Persona Responsable ")
        c.font = Font(name=arial, bold=True, size=10)
        c.alignment = center_wrap
        cellb(r, 3, Border(left=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin')))
        cellb(r, 4, Border(top=Side(style='thin'), bottom=Side(style='thin')))

        c = ws.cell(r, 5, "Albergue")
        c.font = Font(name=arial, bold=True, size=10)
        c.alignment = center_wrap
        cellb(r, 5, thin)

        c = ws.cell(r, 6, "Perfil")
        c.font = Font(name=arial, bold=True, size=10)
        c.alignment = center_wrap
        cellb(r, 6, Border(left=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin')))

        c = ws.cell(r, 7, "Teléfono")
        c.font = Font(name=arial, bold=True, size=10)
        c.alignment = center_wrap
        cellb(r, 7, Border(left=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin')))

        c = ws.cell(r, 8, "Expediente")
        c.font = Font(name=arial, bold=True, size=10)
        c.alignment = center_wrap
        cellb(r, 8, thin)
        r += 1

        # === ROW 7-8: Values ===
        row7 = r
        ws.merge_cells(start_row=r, start_column=1, end_row=r+1, end_column=2)
        c = ws.cell(r, 1, d.get("identidad", ""))
        c.font = Font(name=arial, size=14)
        c.number_format = "0;[Red]0"
        c.alignment = center_wrap
        cellb(r, 1, thin)
        cellb(r, 2, thin_top_bottom)
        cellb(r+1, 1, Border(left=Side(style='thin'), bottom=Side(style='thin')))
        cellb(r+1, 2, Border(right=Side(style='thin'), bottom=Side(style='thin')))

        ws.merge_cells(start_row=r, start_column=3, end_row=r+1, end_column=4)
        c = ws.cell(r, 3, d.get("persona_responsable", ""))
        c.font = Font(name=arial, size=14)
        c.alignment = center_wrap
        cellb(r, 3, thin)
        cellb(r, 4, thin_top_bottom)
        cellb(r+1, 3, Border(left=Side(style='thin'), bottom=Side(style='thin')))
        cellb(r+1, 4, Border(right=Side(style='thin'), bottom=Side(style='thin')))

        ws.merge_cells(start_row=r, start_column=5, end_row=r+1, end_column=5)
        c = ws.cell(r, 5, d.get("albergue", ""))
        c.font = Font(name=arial, size=14)
        c.alignment = center_wrap
        cellb(r, 5, thin)
        cellb(r+1, 5, Border(left=Side(style='thin'), right=Side(style='thin'), bottom=Side(style='thin')))

        ws.merge_cells(start_row=r, start_column=6, end_row=r+1, end_column=6)
        c = ws.cell(r, 6, d.get("perfil", ""))
        c.font = Font(name=arial, size=14)
        c.alignment = center_wrap
        cellb(r, 6, thin)
        cellb(r+1, 6, Border(left=Side(style='thin'), right=Side(style='thin'), bottom=Side(style='thin')))

        ws.merge_cells(start_row=r, start_column=7, end_row=r+1, end_column=7)
        c = ws.cell(r, 7, d.get("telefono", ""))
        c.font = Font(name=arial, size=12)
        c.alignment = center_wrap
        cellb(r, 7, thin)
        cellb(r+1, 7, Border(left=Side(style='thin'), right=Side(style='thin'), bottom=Side(style='thin')))

        ws.merge_cells(start_row=r, start_column=8, end_row=r+1, end_column=8)
        c = ws.cell(r, 8, d.get("expediente", ""))
        c.font = Font(name=arial, size=14)
        c.alignment = center_wrap
        cellb(r, 8, thin)
        cellb(r+1, 8, Border(left=Side(style='thin'), right=Side(style='thin'), bottom=Side(style='thin')))
        r += 2

        # === ROW 9-10: Domicilio ===
        ws.merge_cells(start_row=r, start_column=1, end_row=r+1, end_column=2)
        c = ws.cell(r, 1, "Domicilio del  Paciente:")
        c.font = Font(name=arial, bold=True, size=10)
        c.alignment = center_nowrap
        cellb(r, 1, thin)
        cellb(r, 2, thin_top_bottom)
        cellb(r+1, 1, Border(left=Side(style='thin'), bottom=Side(style='thin')))
        cellb(r+1, 2, Border(right=Side(style='thin'), bottom=Side(style='thin')))

        ws.merge_cells(start_row=r, start_column=3, end_row=r+1, end_column=8)
        domicilio_val = d.get("domicilio", "")
        tel2 = d.get("telefono2", "")
        tel3 = d.get("telefono3", "")
        extras = " | ".join(filter(None, [tel2, tel3]))
        domicilio_final = f"{domicilio_val} | {extras}" if extras else domicilio_val
        c = ws.cell(r, 3, domicilio_final)
        c.font = Font(name=arial, size=12)
        c.alignment = center_nowrap
        cellb(r, 3, thin)
        _apply_borders_range(ws, r, 4, r, 7, thin_top)
        cellb(r, 8, Border(right=Side(style='thin'), top=Side(style='thin')))
        cellb(r+1, 3, Border(left=Side(style='thin'), bottom=Side(style='thin')))
        _apply_borders_range(ws, r+1, 4, r+1, 7, thin_bottom)
        cellb(r+1, 8, Border(right=Side(style='thin'), bottom=Side(style='thin')))
        r += 2

        # === ROW 11: Spacer with top border ===
        _apply_borders_range(ws, r, 1, r, 8, thin_top)
        for cc in range(1, 9):
            ws.cell(r, cc).alignment = Alignment(vertical="center", wrap_text=True)
        r += 1

        # === ROW 12: HEA header ===
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        c = ws.cell(r, 1, "History of Present Illness/Historia de Enfermedad Actual:      ")
        c.font = Font(name=arial, bold=True, size=10)
        c.alignment = center_hnone_wrap
        _apply_borders_range(ws, r, 1, r, 8, thin_top_bottom)
        cellb(r, 1, thin)
        cellb(r, 8, thin)
        r += 1

        # === ROW 13-17: HEA value ===
        ws.merge_cells(start_row=r, start_column=1, end_row=r+4, end_column=8)
        c = ws.cell(r, 1, d.get("historia_enfermedad", ""))
        c.font = Font(name=arial, size=14)
        c.alignment = center_wrap
        cellb(r, 1, thin)
        _apply_borders_range(ws, r, 2, r, 7, thin_top)
        cellb(r, 8, Border(right=Side(style='thin'), top=Side(style='thin')))
        thin_edges(r, r + 4)
        cellb(r+4, 1, Border(left=Side(style='thin'), bottom=Side(style='thin')))
        _apply_borders_range(ws, r+4, 2, r+4, 7, thin_bottom)
        cellb(r+4, 8, Border(right=Side(style='thin'), bottom=Side(style='thin')))
        r += 5

        # === ROW 18: Spacer ===
        c = ws.cell(r, 1)
        c.alignment = vtop
        c.border = Border(left=Side(style='thin'))
        r += 1

        # === ROW 19: Medical History header ===
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        c = ws.cell(r, 1, "Medical History/Antecedentes:")
        c.font = Font(name=arial, bold=True, size=10)
        c.alignment = center_hnone_wrap
        _apply_borders_range(ws, r, 1, r, 8, thin_top_bottom)
        cellb(r, 1, thin)
        cellb(r, 8, thin)
        r += 1

        # === ROW 20-21: Previous illness ===
        ws.merge_cells(start_row=r, start_column=1, end_row=r+1, end_column=3)
        c = ws.cell(r, 1, "Previous illness/ Enfermedades anteriores:")
        c.font = Font(name=arial, bold=True, size=10)
        c.alignment = center_nowrap
        cellb(r, 1, thin)
        cellb(r, 2, thin_top)
        cellb(r, 3, Border(right=Side(style='thin'), top=Side(style='thin')))
        cellb(r+1, 1, Border(left=Side(style='thin'), bottom=Side(style='thin')))
        cellb(r+1, 2, thin_bottom)
        cellb(r+1, 3, Border(right=Side(style='thin'), bottom=Side(style='thin')))

        ws.merge_cells(start_row=r, start_column=4, end_row=r+1, end_column=8)
        c = ws.cell(r, 4, d.get("enfermedades_previas", ""))
        c.font = Font(name=arial, size=14)
        c.alignment = center_wrap
        cellb(r, 4, thin)
        _apply_borders_range(ws, r, 5, r, 7, thin_top)
        cellb(r, 8, Border(right=Side(style='thin'), top=Side(style='thin')))
        cellb(r+1, 4, Border(left=Side(style='thin'), bottom=Side(style='thin')))
        _apply_borders_range(ws, r+1, 5, r+1, 7, thin_bottom)
        cellb(r+1, 8, Border(right=Side(style='thin'), bottom=Side(style='thin')))
        r += 2

        # === ROW 22-23: Past surgeries ===
        ws.merge_cells(start_row=r, start_column=1, end_row=r+1, end_column=3)
        c = ws.cell(r, 1, "Past surgeries/ Cirugías anteriores:")
        c.font = Font(name=arial, bold=True, size=10)
        c.alignment = center_nowrap
        cellb(r, 1, thin)
        cellb(r, 2, thin_top)
        cellb(r, 3, Border(right=Side(style='thin'), top=Side(style='thin')))
        cellb(r+1, 1, Border(left=Side(style='thin'), bottom=Side(style='thin')))
        cellb(r+1, 2, thin_bottom)
        cellb(r+1, 3, Border(right=Side(style='thin'), bottom=Side(style='thin')))

        ws.merge_cells(start_row=r, start_column=4, end_row=r+1, end_column=8)
        c = ws.cell(r, 4, d.get("cirugias_previas", ""))
        c.font = Font(name=arial, size=14)
        c.alignment = center_wrap
        cellb(r, 4, thin)
        _apply_borders_range(ws, r, 5, r, 7, thin_top)
        cellb(r, 8, Border(right=Side(style='thin'), top=Side(style='thin')))
        cellb(r+1, 4, Border(left=Side(style='thin'), bottom=Side(style='thin')))
        _apply_borders_range(ws, r+1, 5, r+1, 7, thin_bottom)
        cellb(r+1, 8, Border(right=Side(style='thin'), bottom=Side(style='thin')))
        r += 2

        # === ROW 24-25: Allergies ===
        ws.merge_cells(start_row=r, start_column=1, end_row=r+1, end_column=3)
        c = ws.cell(r, 1, "Allergies/Alergias:")
        c.font = Font(name=arial, bold=True, size=10)
        c.alignment = center_nowrap
        cellb(r, 1, thin)
        cellb(r, 2, thin_top)
        cellb(r, 3, Border(right=Side(style='thin'), top=Side(style='thin')))
        cellb(r+1, 1, Border(left=Side(style='thin'), bottom=Side(style='thin')))
        cellb(r+1, 2, thin_bottom)
        cellb(r+1, 3, Border(right=Side(style='thin'), bottom=Side(style='thin')))

        ws.merge_cells(start_row=r, start_column=4, end_row=r+1, end_column=8)
        c = ws.cell(r, 4, d.get("alergias", ""))
        c.font = Font(name=arial, size=14)
        c.alignment = center_wrap
        cellb(r, 4, thin)
        _apply_borders_range(ws, r, 5, r, 7, thin_top)
        cellb(r, 8, Border(right=Side(style='thin'), top=Side(style='thin')))
        cellb(r+1, 4, Border(left=Side(style='thin'), bottom=Side(style='thin')))
        _apply_borders_range(ws, r+1, 5, r+1, 7, thin_bottom)
        cellb(r+1, 8, Border(right=Side(style='thin'), bottom=Side(style='thin')))
        r += 2

        # === ROW 26-27: Other ===
        ws.merge_cells(start_row=r, start_column=1, end_row=r+1, end_column=3)
        c = ws.cell(r, 1, "Other/Otros Antecedentes")
        c.font = Font(name=arial, bold=True, size=10)
        c.alignment = center_wrap
        cellb(r, 1, thin)
        cellb(r, 2, thin_top)
        cellb(r, 3, Border(right=Side(style='thin'), top=Side(style='thin')))
        cellb(r+1, 1, Border(left=Side(style='thin'), bottom=Side(style='thin')))
        cellb(r+1, 2, thin_bottom)
        cellb(r+1, 3, Border(right=Side(style='thin'), bottom=Side(style='thin')))

        ws.merge_cells(start_row=r, start_column=4, end_row=r+1, end_column=8)
        c = ws.cell(r, 4, d.get("otros_antecedentes", ""))
        c.font = Font(name=arial, size=14)
        c.alignment = center_wrap
        cellb(r, 4, thin)
        _apply_borders_range(ws, r, 5, r, 7, thin_top)
        cellb(r, 8, Border(right=Side(style='thin'), top=Side(style='thin')))
        cellb(r+1, 4, Border(left=Side(style='thin'), bottom=Side(style='thin')))
        _apply_borders_range(ws, r+1, 5, r+1, 7, thin_bottom)
        cellb(r+1, 8, Border(right=Side(style='thin'), bottom=Side(style='thin')))
        r += 2

        # === ROW 28: blank ===
        r += 1

        # === ROW 29-30: Vital signs ===
        ws.merge_cells(start_row=r, start_column=1, end_row=r+1, end_column=1)
        c = ws.cell(r, 1, "P.A./.B P:")
        c.font = Font(name=arial, bold=True, size=10)
        c.alignment = center_wrap
        cellb(r, 1, thin)
        cellb(r+1, 1, Border(left=Side(style='thin'), right=Side(style='thin'), bottom=Side(style='thin')))

        ws.merge_cells(start_row=r, start_column=2, end_row=r+1, end_column=2)
        c = ws.cell(r, 2, d.get("presion_arterial", ""))
        c.font = Font(name=arial, size=11)
        c.alignment = center_wrap
        cellb(r, 2, thin)
        cellb(r+1, 2, Border(left=Side(style='thin'), right=Side(style='thin'), bottom=Side(style='thin')))

        ws.merge_cells(start_row=r, start_column=3, end_row=r+1, end_column=3)
        fc_val = d.get("fc", "")
        c = ws.cell(r, 3, f"F.C.: {fc_val}x'" if fc_val else "")
        c.font = Font(name=arial, bold=True, size=10)
        c.alignment = center_wrap
        cellb(r, 3, thin)
        cellb(r+1, 3, Border(left=Side(style='thin'), right=Side(style='thin'), bottom=Side(style='thin')))

        pulso_val = d.get("pulso", "")
        c = ws.cell(r, 4, f"Pulso: {pulso_val}x'" if pulso_val else "")
        c.font = Font(name=arial, bold=True, size=9)
        c.alignment = center_wrap
        cellb(r, 4, thin)

        ws.merge_cells(start_row=r, start_column=5, end_row=r+1, end_column=5)
        temp_val = d.get("temperatura", "")
        if temp_val and not temp_val.endswith('°C'):
            temp_str = f"T°: {temp_val}°C"
        else:
            temp_str = f"T°: {temp_val}" if temp_val else ""
        c = ws.cell(r, 5, temp_str)
        c.font = Font(name=arial, bold=True, size=11.5)
        c.alignment = center_wrap
        cellb(r, 5, thin)
        cellb(r+1, 5, Border(left=Side(style='thin'), right=Side(style='thin'), bottom=Side(style='thin')))

        ws.merge_cells(start_row=r, start_column=6, end_row=r+1, end_column=6)
        fr_val = d.get("fr", "")
        c = ws.cell(r, 6, f"F.R.: {fr_val}x'" if fr_val else "")
        c.font = Font(name=arial, bold=True, size=10)
        c.alignment = center_wrap
        cellb(r, 6, Border(right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin')))
        cellb(r+1, 6, Border(right=Side(style='thin'), bottom=Side(style='thin')))

        c = ws.cell(r, 7, "Peso/Weight:")
        c.font = Font(name=arial, bold=True, size=10)
        c.alignment = center_vwrap
        cellb(r, 7, thin)

        peso_val = d.get("peso", "")
        c = ws.cell(r, 8, f"{peso_val} kg" if peso_val and not str(peso_val).strip().endswith("kg") else peso_val if peso_val else "")
        c.font = Font(name=arial, bold=True, size=10)
        c.alignment = left_center_wrap
        cellb(r, 8, thin)

        r += 1

        talla_val = d.get("talla", "")
        c = ws.cell(r, 4, f"Talla: {talla_val}" if talla_val else "")
        c.font = Font(name=arial, bold=True, size=9)
        c.alignment = center_wrap
        cellb(r, 4, thin)

        c = ws.cell(r, 7, "B.M.I.:")
        c.font = Font(name=arial, bold=True, size=10)
        c.alignment = right_center_wrap
        cellb(r, 7, Border(left=Side(style='thin'), right=Side(style='thin'), bottom=Side(style='thin')))

        c = ws.cell(r, 8, d.get("bmi", ""))
        c.font = Font(name=arial, bold=True, size=9)
        c.alignment = left_center_wrap
        cellb(r, 8, Border(left=Side(style='thin'), right=Side(style='thin'), bottom=Side(style='thin')))
        r += 1

        # === ROW 31: Physical Exam header ===
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        c = ws.cell(r, 1, "Physical Exam /Examen Físico ")
        c.font = Font(name=arial, bold=True, size=10)
        c.alignment = center_hnone
        _apply_borders_range(ws, r, 1, r, 8, thin_top_bottom)
        cellb(r, 1, thin)
        cellb(r, 8, thin)
        r += 1

        # === ROW 32-36: Physical exam value ===
        ws.merge_cells(start_row=r, start_column=1, end_row=r+4, end_column=8)
        c = ws.cell(r, 1, d.get("examen_fisico", ""))
        c.font = Font(name=arial, size=12)
        c.alignment = center_wrap
        cellb(r, 1, thin)
        _apply_borders_range(ws, r, 2, r, 7, thin_top)
        cellb(r, 8, Border(right=Side(style='thin'), top=Side(style='thin')))
        thin_edges(r, r + 4)
        cellb(r+4, 1, Border(left=Side(style='thin'), bottom=Side(style='thin')))
        _apply_borders_range(ws, r+4, 2, r+4, 7, thin_bottom)
        cellb(r+4, 8, Border(right=Side(style='thin'), bottom=Side(style='thin')))
        r += 5

        # === ROW 37: Diagnosis header ===
        c = ws.cell(r, 4, "Diagnosis/ Diagnóstico")
        c.font = Font(name=arial, bold=True, size=10)
        ws.cell(r, 5).font = Font(name=arial, bold=True, size=10)
        r += 1

        # === ROW 38-42: Diagnosis value ===
        ws.merge_cells(start_row=r, start_column=1, end_row=r+4, end_column=8)
        c = ws.cell(r, 1, d.get("diagnostico", ""))
        c.font = Font(name=arial, bold=True, size=14)
        c.alignment = center_wrap
        cellb(r, 1, thin)
        _apply_borders_range(ws, r, 2, r, 7, thin_top)
        cellb(r, 8, Border(right=Side(style='thin'), top=Side(style='thin')))
        thin_edges(r, r + 4)
        cellb(r+4, 1, Border(left=Side(style='thin'), bottom=Side(style='thin')))
        _apply_borders_range(ws, r+4, 2, r+4, 7, thin_bottom)
        cellb(r+4, 8, Border(right=Side(style='thin'), bottom=Side(style='thin')))
        r += 5

        # === ROW 43: blank ===
        r += 1

        # === ROW 44: Doctor name ===
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
        c = ws.cell(r, 2, d.get("nombre_medico", ""))
        c.font = Font(name=arial, size=14)
        c.alignment = center_hnone_wrap
        _apply_borders_range(ws, r, 2, r, 7, medium_bottom)
        r += 1

        # === ROW 45: Doctor label ===
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
        c = ws.cell(r, 2, "    Nombre del Médico")
        c.font = Font(name=arial, bold=True, size=10)
        c.alignment = center_hnone
        _apply_borders_range(ws, r, 2, r, 7, medium_top)
        r += 3

        # === ROW 48: Surgeon ===
        c = ws.cell(r, 1, "Surgeon/ Cirujano:")
        c.font = Font(name=arial, bold=True, size=10)
        cellb(r, 1, medium_bottom)
        ws.cell(r, 2).font = Font(name=arial, bold=True, size=10)
        cellb(r, 2, medium_bottom)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=8)
        c = ws.cell(r, 3, d.get("cirujano", ""))
        c.font = Font(name=arial, size=10)
        c.alignment = center_wrap
        _apply_borders_range(ws, r, 3, r, 8, medium_bottom)

        # === ROW 49: spacer (example has bold Arial 10 on A49) ===
        r += 1
        ws.cell(r, 1).font = Font(name=arial, bold=True, size=10)
        r += 1

        # === ROW 50: Surgery Date ===
        label = "Surgery Date/ Day of the Week   Fecha de Cirugía/Día de la Semana:"
        c = ws.cell(r, 1, label)
        c.font = Font(name=arial, bold=True, size=10)
        _apply_borders_range(ws, r, 1, r, 5, medium_bottom)
        for cc in range(2, 6):
            ws.cell(r, cc).font = Font(name=arial, bold=True, size=10)
        ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=8)
        c = ws.cell(r, 6, "")
        c.font = Font(name=arial, size=10)
        c.alignment = center_hnone_wrap
        _apply_borders_range(ws, r, 6, r, 8, medium_bottom)

        # Column widths
        for col_letter, w in col_widths.items():
            ws.column_dimensions[col_letter].width = w
        ws.column_dimensions['I'].bestFit = True

        # Row heights
        for rel_row, h in row_heights.items():
            actual_row = rel_row
            if actual_row <= ws.max_row:
                ws.row_dimensions[actual_row].height = h

        # Shrink to fit: en celdas combinadas Excel ignora shrink_to_fit, por lo que
        # se reduce el tamaño de la fuente hasta que el texto quepa en el área combinada.
        for mr in ws.merged_cells.ranges:
            c = ws.cell(mr.min_row, mr.min_col)
            if c.value is not None:
                base = c.font.size or 10
                fit = _fit_font_size(ws, mr.min_row, mr.min_col, mr.max_row, mr.max_col, c.value, base)
                if fit < base:
                    c.font = Font(name=c.font.name or 'Arial', size=fit, bold=c.font.bold, color=c.font.color)

        # Shrink to fit: si el texto no cabe en la celda, se minimiza hasta que quepa
        for row_cells in ws.iter_rows():
            for cell in row_cells:
                if cell.value is not None and cell.font.bold is not True:
                    a = cell.alignment
                    cell.alignment = Alignment(
                        horizontal=a.horizontal,
                        vertical=a.vertical,
                        wrap_text=True,
                        shrink_to_fit=True,
                    )

    for idx, record in enumerate(records):
        d = record.data if record.data else {}
        data_values = [str(v) for v in d.values() if v not in (None, '')]

        def _is_data_value(s):
            for v in data_values:
                if s == v:
                    return True
                if len(v) >= 3 and v in s:
                    return True
            return False

        ws = wb.create_sheet()
        ws.title = f"Hoja{idx+1}"
        _write_record_sheet(ws, d, styles)
        for row in ws.iter_rows(min_row=1, max_row=52, max_col=11):
            for cell in row:
                if cell.font.name in (None, 'Calibri') and cell.font.size == 11:
                    cell.font = Font(name=arial, size=10)
                if cell.value not in (None, '') and _is_data_value(str(cell.value)):
                    a = cell.alignment
                    cell.alignment = Alignment(horizontal=a.horizontal, vertical=a.vertical,
                                               wrap_text=a.wrap_text, shrink_to_fit=True)
                b = cell.border
                new_sides = {}
                for side_name in ('left', 'right', 'top', 'bottom'):
                    side = getattr(b, side_name)
                    if side is not None and side.style and (side.color is None or side.color.type not in ('auto',)):
                        new_sides[side_name] = Side(style=side.style, color=Color(auto=True))
                if new_sides:
                    cell.border = Border(**{**{
                        'left': b.left, 'right': b.right, 'top': b.top, 'bottom': b.bottom,
                    }, **new_sides})

    wb.remove(default_sheet)
    wb.save(filepath)


def _fit_font_size(ws, r1, c1, r2, c2, text, base_size):
    """Calcula el mayor tamaño de fuente (Arial) que permite que `text` quepa dentro
    del área combinada (r1,c1)-(r2,c2), considerando el ancho de las columnas y la
    altura de las filas de la hoja."""
    if not text:
        return base_size

    total_width = 0.0
    for cc in range(c1, c2 + 1):
        w = ws.column_dimensions[get_column_letter(cc)].width
        total_width += w if w else 8.43

    total_height = 0.0
    for rr in range(r1, r2 + 1):
        h = ws.row_dimensions[rr].height
        total_height += h if h else 15.0

    size = float(base_size)
    min_size = 5.0
    while size > min_size:
        # Caracteres aproximados por línea: ancho total (unidades ~ caracteres) por
        # proporción de la fuente base (10pt) respecto al tamaño probado.
        chars_per_line = max(1, int(total_width * (10.0 / size) * 0.92))
        lines_needed = 0
        for segment in str(text).split("\n"):
            lines_needed += max(1, math.ceil(len(segment) / chars_per_line))
        line_height = size * 1.35  # altura aproximada de línea en puntos
        if lines_needed * line_height <= total_height:
            return size
        size -= 0.5
    return min_size


def _apply_border(ws, row, col, border):
    ws.cell(row, col).border = border


def _apply_borders_range(ws, r1, c1, r2, c2, border):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(r, c).border = border
